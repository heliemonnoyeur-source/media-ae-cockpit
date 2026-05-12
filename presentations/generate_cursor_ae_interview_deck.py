from pathlib import Path

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_AUTO_SHAPE_TYPE
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.util import Inches, Pt


OUTPUT_FILE = Path(__file__).resolve().parent / "Cursor_AE_Interview_Deck.pptx"

PROFILE = {
    "candidate_name": "Your Name",
    "role_title": "Account Executive Candidate",
    "contact_line": "you@example.com | linkedin.com/in/yourname | City, ST",
    "north_star": (
        "Help Cursor turn bottom-up enthusiasm into repeatable, multi-threaded"
        " enterprise revenue."
    ),
    "top_results": [
        "Replace with a headline result, e.g. 142% of quota across strategic accounts.",
        "Replace with a customer story, e.g. grew a pilot into a 6-figure expansion.",
        "Replace with a GTM example, e.g. built outbound plays for a new product line.",
    ],
}

PALETTE = {
    "bg": RGBColor(13, 17, 23),
    "card": RGBColor(22, 27, 34),
    "card_alt": RGBColor(24, 34, 46),
    "accent": RGBColor(94, 234, 212),
    "accent_alt": RGBColor(96, 165, 250),
    "text": RGBColor(248, 250, 252),
    "muted": RGBColor(148, 163, 184),
    "success": RGBColor(74, 222, 128),
    "warning": RGBColor(251, 191, 36),
}

FONT = {
    "title": "Aptos Display",
    "body": "Aptos",
}


def set_bg(slide, color):
    fill = slide.background.fill
    fill.solid()
    fill.fore_color.rgb = color


def add_textbox(
    slide,
    left,
    top,
    width,
    height,
    text,
    font_size,
    color,
    bold=False,
    font_name=None,
    align=PP_ALIGN.LEFT,
    valign=MSO_ANCHOR.TOP,
):
    shape = slide.shapes.add_textbox(left, top, width, height)
    frame = shape.text_frame
    frame.word_wrap = True
    frame.vertical_anchor = valign
    paragraph = frame.paragraphs[0]
    paragraph.alignment = align
    run = paragraph.add_run()
    run.text = text
    run.font.size = Pt(font_size)
    run.font.bold = bold
    run.font.color.rgb = color
    run.font.name = font_name or FONT["body"]
    return shape


def add_title_block(slide, title, subtitle=None, eyebrow=None):
    if eyebrow:
        add_textbox(
            slide,
            Inches(0.8),
            Inches(0.45),
            Inches(4.5),
            Inches(0.4),
            eyebrow.upper(),
            12,
            PALETTE["accent"],
            bold=True,
        )
    add_textbox(
        slide,
        Inches(0.8),
        Inches(0.9),
        Inches(11.2),
        Inches(0.9),
        title,
        26,
        PALETTE["text"],
        bold=True,
        font_name=FONT["title"],
    )
    if subtitle:
        add_textbox(
            slide,
            Inches(0.8),
            Inches(1.65),
            Inches(11.0),
            Inches(0.8),
            subtitle,
            13,
            PALETTE["muted"],
        )


def add_footer(slide, page_number):
    add_textbox(
        slide,
        Inches(0.8),
        Inches(7.0),
        Inches(5.5),
        Inches(0.3),
        "Cursor AE interview deck",
        9,
        PALETTE["muted"],
    )
    add_textbox(
        slide,
        Inches(12.0),
        Inches(7.0),
        Inches(0.5),
        Inches(0.3),
        str(page_number),
        9,
        PALETTE["muted"],
        align=PP_ALIGN.RIGHT,
    )


def add_card(slide, left, top, width, height, title, bullets, fill_color):
    shape = slide.shapes.add_shape(
        MSO_AUTO_SHAPE_TYPE.ROUNDED_RECTANGLE, left, top, width, height
    )
    shape.fill.solid()
    shape.fill.fore_color.rgb = fill_color
    shape.line.color.rgb = fill_color
    shape.line.width = Pt(1)

    add_textbox(
        slide,
        left + Inches(0.2),
        top + Inches(0.18),
        width - Inches(0.4),
        Inches(0.35),
        title,
        15,
        PALETTE["text"],
        bold=True,
    )

    body = slide.shapes.add_textbox(
        left + Inches(0.2), top + Inches(0.6), width - Inches(0.4), height - Inches(0.75)
    )
    frame = body.text_frame
    frame.word_wrap = True
    for index, bullet in enumerate(bullets):
        paragraph = frame.paragraphs[0] if index == 0 else frame.add_paragraph()
        paragraph.text = bullet
        paragraph.level = 0
        paragraph.bullet = True
        paragraph.font.size = Pt(11)
        paragraph.font.color.rgb = PALETTE["text"]
        paragraph.font.name = FONT["body"]


def add_bullet_column(slide, left, top, width, title, bullets):
    add_textbox(slide, left, top, width, Inches(0.3), title, 15, PALETTE["accent"], bold=True)
    body = slide.shapes.add_textbox(left, top + Inches(0.45), width, Inches(4.6))
    frame = body.text_frame
    frame.word_wrap = True
    for index, bullet in enumerate(bullets):
        paragraph = frame.paragraphs[0] if index == 0 else frame.add_paragraph()
        paragraph.text = bullet
        paragraph.level = 0
        paragraph.bullet = True
        paragraph.font.size = Pt(14)
        paragraph.font.color.rgb = PALETTE["text"]
        paragraph.font.name = FONT["body"]


def add_score_row(slide, top, label, value, detail, color):
    shape = slide.shapes.add_shape(
        MSO_AUTO_SHAPE_TYPE.ROUNDED_RECTANGLE,
        Inches(0.85),
        top,
        Inches(11.6),
        Inches(0.72),
    )
    shape.fill.solid()
    shape.fill.fore_color.rgb = PALETTE["card"]
    shape.line.color.rgb = PALETTE["card"]

    tag = slide.shapes.add_shape(
        MSO_AUTO_SHAPE_TYPE.ROUNDED_RECTANGLE,
        Inches(0.95),
        top + Inches(0.12),
        Inches(1.7),
        Inches(0.45),
    )
    tag.fill.solid()
    tag.fill.fore_color.rgb = color
    tag.line.color.rgb = color

    add_textbox(
        slide,
        Inches(1.02),
        top + Inches(0.18),
        Inches(1.55),
        Inches(0.2),
        label,
        12,
        PALETTE["bg"],
        bold=True,
    )
    add_textbox(
        slide,
        Inches(2.95),
        top + Inches(0.12),
        Inches(1.45),
        Inches(0.28),
        value,
        18,
        PALETTE["text"],
        bold=True,
    )
    add_textbox(
        slide,
        Inches(4.35),
        top + Inches(0.16),
        Inches(7.7),
        Inches(0.25),
        detail,
        11,
        PALETTE["muted"],
    )


def build_deck():
    presentation = Presentation()
    presentation.slide_width = Inches(13.333)
    presentation.slide_height = Inches(7.5)
    blank = presentation.slide_layouts[6]

    slides = []

    slide = presentation.slides.add_slide(blank)
    set_bg(slide, PALETTE["bg"])
    add_textbox(
        slide,
        Inches(0.8),
        Inches(0.65),
        Inches(3.5),
        Inches(0.35),
        "CURSOR INTERVIEW PRESENTATION",
        12,
        PALETTE["accent"],
        bold=True,
    )
    add_textbox(
        slide,
        Inches(0.8),
        Inches(1.2),
        Inches(8.8),
        Inches(1.1),
        "How I'd Win as an Account Executive at Cursor",
        28,
        PALETTE["text"],
        bold=True,
        font_name=FONT["title"],
    )
    add_textbox(
        slide,
        Inches(0.8),
        Inches(2.4),
        Inches(7.0),
        Inches(0.8),
        PROFILE["north_star"],
        17,
        PALETTE["muted"],
    )
    badge = slide.shapes.add_shape(
        MSO_AUTO_SHAPE_TYPE.ROUNDED_RECTANGLE,
        Inches(0.8),
        Inches(5.6),
        Inches(4.8),
        Inches(0.85),
    )
    badge.fill.solid()
    badge.fill.fore_color.rgb = PALETTE["card_alt"]
    badge.line.color.rgb = PALETTE["card_alt"]
    add_textbox(
        slide,
        Inches(1.05),
        Inches(5.84),
        Inches(4.2),
        Inches(0.3),
        f"{PROFILE['candidate_name']} | {PROFILE['role_title']}",
        16,
        PALETTE["text"],
        bold=True,
    )
    add_textbox(
        slide,
        Inches(1.05),
        Inches(6.18),
        Inches(4.1),
        Inches(0.2),
        PROFILE["contact_line"],
        10,
        PALETTE["muted"],
    )
    accent = slide.shapes.add_shape(
        MSO_AUTO_SHAPE_TYPE.ROUNDED_RECTANGLE,
        Inches(10.2),
        Inches(1.3),
        Inches(2.2),
        Inches(4.7),
    )
    accent.fill.solid()
    accent.fill.fore_color.rgb = PALETTE["card"]
    accent.line.color.rgb = PALETTE["card"]
    add_textbox(
        slide,
        Inches(10.55),
        Inches(1.7),
        Inches(1.6),
        Inches(0.4),
        "AE",
        36,
        PALETTE["accent"],
        bold=True,
        align=PP_ALIGN.CENTER,
    )
    add_textbox(
        slide,
        Inches(10.45),
        Inches(2.45),
        Inches(1.8),
        Inches(2.4),
        "Revenue\noperator\nfor an AI-\nnative era",
        18,
        PALETTE["text"],
        bold=True,
        align=PP_ALIGN.CENTER,
        valign=MSO_ANCHOR.MIDDLE,
    )
    slides.append(slide)

    slide = presentation.slides.add_slide(blank)
    set_bg(slide, PALETTE["bg"])
    add_title_block(
        slide,
        "My thesis for the role",
        "Cursor has the ingredients for fast growth: product love, clear workflow value, and an urgent budget conversation around AI tooling.",
        eyebrow="Executive summary",
    )
    add_card(
        slide,
        Inches(0.8),
        Inches(2.1),
        Inches(3.9),
        Inches(3.9),
        "1. Win where product pull already exists",
        [
            "Prioritize accounts where engineers are self-educating on AI coding tools.",
            "Turn bottom-up usage into a business case around velocity, quality, and security.",
            "Use proof, not hype: demo measurable team outcomes quickly.",
        ],
        PALETTE["card"],
    )
    add_card(
        slide,
        Inches(4.95),
        Inches(2.1),
        Inches(3.9),
        Inches(3.9),
        "2. Sell a rollout, not just seats",
        [
            "Map each deal to a rollout plan with champions, leaders, and IT/security stakeholders.",
            "Anchor on adoption milestones that justify expansion and renewal.",
            "Make it easy for managers to see who is getting value and where to scale.",
        ],
        PALETTE["card_alt"],
    )
    add_card(
        slide,
        Inches(9.1),
        Inches(2.1),
        Inches(3.4),
        Inches(3.9),
        "3. Build repeatable expansion plays",
        [
            "Package customer stories by use case, persona, and technical maturity.",
            "Create repeatable paths from one team to many teams.",
            "Operate tightly with product, success, and leadership feedback loops.",
        ],
        PALETTE["card"],
    )
    slides.append(slide)

    slide = presentation.slides.add_slide(blank)
    set_bg(slide, PALETTE["bg"])
    add_title_block(
        slide,
        "Why Cursor, why now",
        "My read on the market conditions that make this role compelling.",
        eyebrow="Market view",
    )
    add_bullet_column(
        slide,
        Inches(0.8),
        Inches(2.0),
        Inches(3.75),
        "What buyers care about",
        [
            "Engineering leaders need concrete productivity gains, not generic AI narratives.",
            "Teams want tools developers actually love enough to use every day.",
            "Security and governance must be part of the rollout conversation early.",
        ],
    )
    add_bullet_column(
        slide,
        Inches(4.75),
        Inches(2.0),
        Inches(3.65),
        "Why Cursor is well positioned",
        [
            "The product sits directly inside the developer workflow where habits are formed.",
            "The value story is intuitive: faster execution, better context, less toil.",
            "Product enthusiasm creates a strong wedge for account expansion.",
        ],
    )
    add_bullet_column(
        slide,
        Inches(8.65),
        Inches(2.0),
        Inches(3.85),
        "Where an AE can add leverage",
        [
            "Translate product pull into executive language around throughput and risk reduction.",
            "Create stakeholder alignment so usage becomes standardized spend.",
            "Build a disciplined operating cadence across sourcing, pilots, and expansions.",
        ],
    )
    slides.append(slide)

    slide = presentation.slides.add_slide(blank)
    set_bg(slide, PALETTE["bg"])
    add_title_block(
        slide,
        "ICP and territory priorities",
        "I would focus where urgency, workflow pain, and internal champions are most likely to converge.",
        eyebrow="Segmentation",
    )
    add_card(
        slide,
        Inches(0.8),
        Inches(2.0),
        Inches(3.8),
        Inches(4.1),
        "Priority A: engineering-led scaleups",
        [
            "Fast-moving software organizations with developer density.",
            "VP Eng / CTO care about shipping speed and retaining strong talent.",
            "Great fit for fast pilots and high-velocity expansions.",
        ],
        PALETTE["card"],
    )
    add_card(
        slide,
        Inches(4.8),
        Inches(2.0),
        Inches(3.8),
        Inches(4.1),
        "Priority B: enterprise innovation teams",
        [
            "Business units already experimenting with AI-enabled development.",
            "Need governance, stakeholder coordination, and a credible rollout path.",
            "High upside if one successful team becomes the internal reference account.",
        ],
        PALETTE["card_alt"],
    )
    add_card(
        slide,
        Inches(8.8),
        Inches(2.0),
        Inches(3.7),
        Inches(4.1),
        "Priority C: expansion inside active users",
        [
            "Accounts where product usage or inbound interest is already visible.",
            "Best route to efficient revenue: prove value, standardize, then scale.",
            "Requires tight discovery, stakeholder mapping, and customer storytelling.",
        ],
        PALETTE["card"],
    )
    slides.append(slide)

    slide = presentation.slides.add_slide(blank)
    set_bg(slide, PALETTE["bg"])
    add_title_block(
        slide,
        "My pipeline generation engine",
        "A balanced approach that mixes signal-based prospecting with customer-led storytelling.",
        eyebrow="Pipeline",
    )
    add_card(
        slide,
        Inches(0.8),
        Inches(2.0),
        Inches(2.9),
        Inches(3.9),
        "Product signals",
        [
            "Prioritize accounts showing developer curiosity or grassroots adoption.",
            "Use these signals to make outreach timely and specific.",
        ],
        PALETTE["card"],
    )
    add_card(
        slide,
        Inches(3.95),
        Inches(2.0),
        Inches(2.9),
        Inches(3.9),
        "Persona-led outreach",
        [
            "Different message for VP Eng, DevEx, security, and team leaders.",
            "Lead with workflow outcomes each persona can own.",
        ],
        PALETTE["card_alt"],
    )
    add_card(
        slide,
        Inches(7.1),
        Inches(2.0),
        Inches(2.9),
        Inches(3.9),
        "Social proof",
        [
            "Package success stories by company type and use case.",
            "Turn customer outcomes into credibility for each next meeting.",
        ],
        PALETTE["card"],
    )
    add_card(
        slide,
        Inches(10.25),
        Inches(2.0),
        Inches(2.1),
        Inches(3.9),
        "Discipline",
        [
            "Weekly review of target accounts, meeting quality, pipeline stage health, and next steps.",
            "High standards on follow-up and multi-threading.",
        ],
        PALETTE["card_alt"],
    )
    slides.append(slide)

    slide = presentation.slides.add_slide(blank)
    set_bg(slide, PALETTE["bg"])
    add_title_block(
        slide,
        "Discovery -> POV -> rollout",
        "The sales motion I would run to convert excitement into durable adoption.",
        eyebrow="Sales process",
    )
    add_bullet_column(
        slide,
        Inches(0.8),
        Inches(2.0),
        Inches(5.7),
        "Deal motion",
        [
            "1. Discovery: isolate high-friction engineering workflows and current tooling gaps.",
            "2. Success criteria: agree on what a strong pilot outcome looks like before it starts.",
            "3. POV: make the product prove value inside a real team workflow, not a synthetic demo.",
            "4. Rollout plan: define who expands usage, who approves spend, and how adoption will be measured.",
            "5. Expansion: reuse the win story to open adjacent teams and higher-level sponsors.",
        ],
    )
    add_card(
        slide,
        Inches(7.0),
        Inches(2.1),
        Inches(5.4),
        Inches(3.9),
        "Stakeholders I want mapped early",
        [
            "Champion: staff engineer / eng manager who can validate workflow improvement.",
            "Economic buyer: VP Eng, CTO, or budget owner tied to productivity.",
            "Security / IT: needed for procurement confidence and rollout speed.",
            "Executive sponsor: leader who can help standardize usage across teams.",
        ],
        PALETTE["card"],
    )
    slides.append(slide)

    slide = presentation.slides.add_slide(blank)
    set_bg(slide, PALETTE["bg"])
    add_title_block(
        slide,
        "Land, prove value, expand",
        "My account planning lens for turning one win into a larger revenue story.",
        eyebrow="Expansion",
    )
    add_card(
        slide,
        Inches(0.8),
        Inches(2.1),
        Inches(3.75),
        Inches(3.9),
        "Land",
        [
            "Start with a motivated team and a crisp business problem.",
            "Keep scope small enough to move fast, but visible enough to matter.",
            "Document the internal narrative before the pilot begins.",
        ],
        PALETTE["card"],
    )
    add_card(
        slide,
        Inches(4.8),
        Inches(2.1),
        Inches(3.75),
        Inches(3.9),
        "Prove",
        [
            "Tie user feedback to concrete workflow improvements and adoption patterns.",
            "Create a simple readout the champion can carry upward.",
            "Show why the product belongs in the standard toolchain.",
        ],
        PALETTE["card_alt"],
    )
    add_card(
        slide,
        Inches(8.8),
        Inches(2.1),
        Inches(3.75),
        Inches(3.9),
        "Expand",
        [
            "Sequence adjacent teams by similarity of pain and leadership structure.",
            "Use early champions to recruit second-order champions.",
            "Protect renewals by staying close to realized value, not shelfware risk.",
        ],
        PALETTE["card"],
    )
    slides.append(slide)

    slide = presentation.slides.add_slide(blank)
    set_bg(slide, PALETTE["bg"])
    add_title_block(
        slide,
        "30 / 60 / 90 day plan",
        "The first quarter should balance product fluency, target account focus, and execution discipline.",
        eyebrow="Ramp plan",
    )
    add_card(
        slide,
        Inches(0.8),
        Inches(2.05),
        Inches(3.8),
        Inches(4.15),
        "Days 1-30: learn and calibrate",
        [
            "Deepen product fluency and customer language.",
            "Review won/lost deals and listen for repeatable patterns.",
            "Build a target account list and first-pass stakeholder hypotheses.",
        ],
        PALETTE["card"],
    )
    add_card(
        slide,
        Inches(4.8),
        Inches(2.05),
        Inches(3.8),
        Inches(4.15),
        "Days 31-60: build pipeline",
        [
            "Launch account-based outreach with crisp persona messaging.",
            "Convert early meetings into proof-oriented next steps.",
            "Pressure test discovery and pilot motions with manager feedback.",
        ],
        PALETTE["card_alt"],
    )
    add_card(
        slide,
        Inches(8.8),
        Inches(2.05),
        Inches(3.75),
        Inches(4.15),
        "Days 61-90: create momentum",
        [
            "Advance qualified opportunities with clear mutual action plans.",
            "Establish expansion paths inside the strongest early accounts.",
            "Share what is working across messaging, personas, and customer objections.",
        ],
        PALETTE["card"],
    )
    slides.append(slide)

    slide = presentation.slides.add_slide(blank)
    set_bg(slide, PALETTE["bg"])
    add_title_block(
        slide,
        "Operating cadence and scorecard",
        "The numbers I would watch to keep quality high and surprises low.",
        eyebrow="Execution",
    )
    add_score_row(
        slide,
        Inches(2.0),
        "Coverage",
        "3-4x",
        "Pipeline coverage by quarter, segmented by new logos, pilot expansions, and strategic bets.",
        PALETTE["accent"],
    )
    add_score_row(
        slide,
        Inches(2.95),
        "Velocity",
        "Tracked weekly",
        "Stage progression, next-step hygiene, and time-to-pilot / time-to-close.",
        PALETTE["accent_alt"],
    )
    add_score_row(
        slide,
        Inches(3.9),
        "Adoption",
        "Core leading signal",
        "Pilot engagement, champion energy, and evidence the product is entering daily workflow.",
        PALETTE["success"],
    )
    add_score_row(
        slide,
        Inches(4.85),
        "Risk",
        "No silent deals",
        "Every meaningful opportunity has an identified blocker, owner, and date to resolve it.",
        PALETTE["warning"],
    )
    add_textbox(
        slide,
        Inches(0.9),
        Inches(6.2),
        Inches(11.2),
        Inches(0.45),
        "Operating principle: tight inspection without micromanaging the customer journey.",
        14,
        PALETTE["text"],
        bold=True,
    )
    slides.append(slide)

    slide = presentation.slides.add_slide(blank)
    set_bg(slide, PALETTE["bg"])
    add_title_block(
        slide,
        "Handling objections and competitive pressure",
        "I prefer to de-risk decisions by making the evaluation concrete and customer-specific.",
        eyebrow="Positioning",
    )
    add_card(
        slide,
        Inches(0.8),
        Inches(2.0),
        Inches(5.7),
        Inches(4.25),
        "Common objections",
        [
            "\"We already have another AI coding tool.\"",
            "\"Security and governance are not ready yet.\"",
            "\"I like it, but I cannot justify broader rollout today.\"",
            "\"Developers may test it, but that does not mean the business should buy it.\"",
        ],
        PALETTE["card"],
    )
    add_card(
        slide,
        Inches(6.85),
        Inches(2.0),
        Inches(5.7),
        Inches(4.25),
        "My response pattern",
        [
            "Clarify the actual job to be done and where the current workflow still breaks.",
            "Anchor on evaluation criteria tied to team outcomes, not feature checklists alone.",
            "Bring in the right stakeholder early so procurement or security does not become a late surprise.",
            "Offer a rollout path that is easy to approve because the proof already exists.",
        ],
        PALETTE["card_alt"],
    )
    slides.append(slide)

    slide = presentation.slides.add_slide(blank)
    set_bg(slide, PALETTE["bg"])
    add_title_block(
        slide,
        "Why me for this role",
        "Replace the placeholders below with your strongest, most relevant proof points.",
        eyebrow="Close",
    )
    add_card(
        slide,
        Inches(0.8),
        Inches(2.0),
        Inches(5.4),
        Inches(4.35),
        "Proof points to customize",
        PROFILE["top_results"],
        PALETTE["card"],
    )
    add_card(
        slide,
        Inches(6.55),
        Inches(2.0),
        Inches(5.95),
        Inches(4.35),
        "Closing statement",
        [
            "I bring a disciplined sales process and enough curiosity to earn credibility with technical buyers.",
            "I know how to turn product excitement into stakeholder alignment and revenue momentum.",
            "If I joined Cursor, my goal would be to help build a category-defining sales motion around an exceptional product.",
        ],
        PALETTE["card_alt"],
    )
    slides.append(slide)

    slide = presentation.slides.add_slide(blank)
    set_bg(slide, PALETTE["bg"])
    add_title_block(
        slide,
        "Questions I'd ask in the interview",
        "Useful if they want the presentation to feel collaborative rather than one-way.",
        eyebrow="Appendix",
    )
    add_bullet_column(
        slide,
        Inches(0.8),
        Inches(2.0),
        Inches(11.5),
        "Sample questions",
        [
            "Which customer profiles are expanding fastest today, and what do those wins have in common?",
            "Where do deals tend to stall most often: technical validation, security, procurement, or stakeholder alignment?",
            "How do the best AEs at Cursor work with product and customer success to accelerate rollout quality?",
            "What leading signals tell you an account is likely to become a durable multi-team customer?",
            "What do you want this hire to prove in the first few months that would feel like a clear win?",
        ],
    )
    slides.append(slide)

    for index, built_slide in enumerate(slides, start=1):
        add_footer(built_slide, index)

    presentation.save(OUTPUT_FILE)


if __name__ == "__main__":
    OUTPUT_FILE.parent.mkdir(parents=True, exist_ok=True)
    build_deck()
    print(f"Created {OUTPUT_FILE}")
