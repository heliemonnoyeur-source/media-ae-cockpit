from pathlib import Path

from pptx import Presentation
from pptx.enum.shapes import MSO_SHAPE_TYPE
from pptx.enum.text import PP_ALIGN
from reportlab.lib.colors import Color
from reportlab.lib.utils import simpleSplit
from reportlab.pdfgen import canvas


EMU_PER_INCH = 914400
POINTS_PER_INCH = 72
NS = {"a": "http://schemas.openxmlformats.org/drawingml/2006/main"}

PPTX_PATH = Path(__file__).resolve().parent / "Cursor_AE_Interview_Deck.pptx"
PDF_PATH = Path(__file__).resolve().parent / "Cursor_AE_Interview_Deck.pdf"


def emu_to_points(value):
    return value * POINTS_PER_INCH / EMU_PER_INCH


def rgb_to_color(rgb, default):
    if rgb is None:
        return default
    return Color(rgb[0] / 255, rgb[1] / 255, rgb[2] / 255)


def safe_rgb(color_format):
    try:
        return color_format.rgb
    except Exception:
        return None


def safe_fill_rgb(fill):
    try:
        return safe_rgb(fill.fore_color)
    except Exception:
        return None


def safe_line_rgb(line):
    try:
        return safe_rgb(line.color)
    except Exception:
        return None


def get_run_or_paragraph_font(paragraph):
    if paragraph.runs:
        return paragraph.runs[0].font
    return paragraph.font


def paragraph_font_size(paragraph, default=12):
    font = get_run_or_paragraph_font(paragraph)
    if font.size is not None:
        return font.size.pt

    ppr = paragraph._p.find("a:pPr", NS)
    if ppr is not None:
        defrpr = ppr.find("a:defRPr", NS)
        if defrpr is not None and defrpr.get("sz"):
            return int(defrpr.get("sz")) / 100

    run = paragraph._p.find("a:r", NS)
    if run is not None:
        rpr = run.find("a:rPr", NS)
        if rpr is not None and rpr.get("sz"):
            return int(rpr.get("sz")) / 100

    return default


def paragraph_bold(paragraph):
    font = get_run_or_paragraph_font(paragraph)
    if font.bold is not None:
        return font.bold

    run = paragraph._p.find("a:r", NS)
    if run is not None:
        rpr = run.find("a:rPr", NS)
        if rpr is not None and rpr.get("b") is not None:
            return rpr.get("b") == "1"
    return False


def paragraph_color(paragraph, default):
    font = get_run_or_paragraph_font(paragraph)
    if font.color is not None:
        rgb = safe_rgb(font.color)
        if rgb is not None:
            return rgb_to_color(rgb, default)

    ppr = paragraph._p.find("a:pPr", NS)
    if ppr is not None:
        defrpr = ppr.find("a:defRPr", NS)
        if defrpr is not None:
            solid = defrpr.find("a:solidFill/a:srgbClr", NS)
            if solid is not None and solid.get("val"):
                value = solid.get("val")
                rgb = tuple(int(value[i : i + 2], 16) for i in (0, 2, 4))
                return rgb_to_color(rgb, default)

    run = paragraph._p.find("a:r", NS)
    if run is not None:
        rpr = run.find("a:rPr", NS)
        if rpr is not None:
            solid = rpr.find("a:solidFill/a:srgbClr", NS)
            if solid is not None and solid.get("val"):
                value = solid.get("val")
                rgb = tuple(int(value[i : i + 2], 16) for i in (0, 2, 4))
                return rgb_to_color(rgb, default)

    return default


def draw_shape_background(pdf, shape, page_height):
    if shape.shape_type != MSO_SHAPE_TYPE.AUTO_SHAPE:
        return

    fill_rgb = safe_fill_rgb(shape.fill)
    if fill_rgb is None:
        return

    x = emu_to_points(shape.left)
    y = page_height - emu_to_points(shape.top + shape.height)
    width = emu_to_points(shape.width)
    height = emu_to_points(shape.height)

    stroke_rgb = safe_line_rgb(shape.line)
    pdf.setFillColor(rgb_to_color(fill_rgb, Color(0, 0, 0)))
    if stroke_rgb is not None:
        pdf.setStrokeColor(rgb_to_color(stroke_rgb, Color(0, 0, 0)))
        stroke = 1
    else:
        stroke = 0

    radius = min(width, height) * 0.08
    pdf.roundRect(x, y, width, height, radius, fill=1, stroke=stroke)


def paragraph_prefix(paragraph, index, paragraph_count):
    if paragraph_count <= 1:
        return ""

    size = paragraph_font_size(paragraph)
    if size >= 18:
        return ""

    return "• "


def draw_text_frame(pdf, shape, page_height):
    if not getattr(shape, "has_text_frame", False):
        return

    paragraphs = [p for p in shape.text_frame.paragraphs if p.text.strip()]
    if not paragraphs:
        return

    x = emu_to_points(shape.left)
    y_top = page_height - emu_to_points(shape.top)
    width = emu_to_points(shape.width)
    height = emu_to_points(shape.height)
    cursor_y = y_top - 2

    for index, paragraph in enumerate(paragraphs):
        text = f"{paragraph_prefix(paragraph, index, len(paragraphs))}{paragraph.text.strip()}"
        font_size = paragraph_font_size(paragraph)
        font_name = "Helvetica-Bold" if paragraph_bold(paragraph) else "Helvetica"
        color = paragraph_color(paragraph, Color(0.97, 0.98, 0.99))
        align = paragraph.alignment
        wrapped = simpleSplit(text, font_name, font_size, max(width - 4, 40))
        line_height = font_size * 1.2

        for line in wrapped:
            if cursor_y - line_height < y_top - height:
                return
            pdf.setFont(font_name, font_size)
            pdf.setFillColor(color)

            if align == PP_ALIGN.CENTER:
                pdf.drawCentredString(x + width / 2, cursor_y - font_size, line)
            elif align == PP_ALIGN.RIGHT:
                pdf.drawRightString(x + width, cursor_y - font_size, line)
            else:
                pdf.drawString(x, cursor_y - font_size, line)

            cursor_y -= line_height

        cursor_y -= font_size * 0.35


def export_pdf():
    presentation = Presentation(PPTX_PATH)
    page_width = emu_to_points(presentation.slide_width)
    page_height = emu_to_points(presentation.slide_height)

    pdf = canvas.Canvas(str(PDF_PATH), pagesize=(page_width, page_height))

    for slide in presentation.slides:
        background = safe_fill_rgb(slide.background.fill)
        pdf.setFillColor(rgb_to_color(background, Color(13 / 255, 17 / 255, 23 / 255)))
        pdf.rect(0, 0, page_width, page_height, fill=1, stroke=0)

        for shape in slide.shapes:
            draw_shape_background(pdf, shape, page_height)

        for shape in slide.shapes:
            draw_text_frame(pdf, shape, page_height)

        pdf.showPage()

    pdf.save()


if __name__ == "__main__":
    export_pdf()
    print(f"Created {PDF_PATH}")
